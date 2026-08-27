# -*- coding: utf-8 -*-
"""Build CBEB-formatted English DOCX from original template + translations, then export PDF."""
import json
import os
import shutil
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = r"C:\Users\GonLu\Downloads\[ACNE] - Artigo CBEB 2026 - V-Final.docx"
TRANS_JSON = os.path.join(ROOT, "papers", "cbeb", "translations_en.json")
FIGS = os.path.join(ROOT, "papers", "sibgrapi", "figs")
OUT_DOCX = os.path.join(ROOT, "papers", "cbeb", "AcneNet_CBEB2026_English_CBEB.docx")
OUT_PDF = os.path.join(ROOT, "papers", "cbeb", "AcneNet_CBEB2026_English_CBEB.pdf")

TABLE_I = [
    ["Characteristic", "Acne04", "Acne Level"],
    ["Objective", "Acne severity classification", "Acne severity classification"],
    ["Number of images", "1,406", "1,457"],
    ["Classes", "4", "4"],
    ["Severity scale", "Hayashi (grades 0–3)", "Hayashi (grades 0–3)"],
    ["Image type", "Frontal face", "Frontal face"],
    ["Organization", "Original corpus", "Folder-organized by severity"],
    ["Use in this work", "Merged into unified pipeline", "Merged into unified pipeline"],
]
TABLE_II = [
    ["Parameter", "Value"],
    ["Framework", "PyTorch"],
    ["Backbone", "ResNet-18 (trained from scratch, weights=None)"],
    ["Optimizer", "Adam"],
    ["Loss function", "Weighted Cross-Entropy"],
    ["Learning rate", "0.0001"],
    ["Scheduler", "ReduceLROnPlateau"],
    ["Batch size", "16"],
    ["Maximum epochs", "40"],
    ["Data augmentation", "Horizontal flip (p=0.5), rotation (±25°), color jitter, Gaussian blur, random erasing"],
    ["Regularization", "Batch Normalization and Dropout"],
    ["Hardware", "NVIDIA GeForce RTX 3060 (12 GB)"],
    ["Model weights", "best_robust_model.pth"],
    ["Inference script", "training/predict.py"],
]
TABLE_III = [
    ["Class", "Precision", "Recall", "F1-score", "Support"],
    ["Grade 0 (Clear)", "0.96", "0.95", "0.95", "318"],
    ["Grade 1 (Mild)", "0.94", "0.94", "0.94", "412"],
    ["Grade 2 (Moderate)", "0.88", "0.92", "0.90", "110"],
    ["Grade 3 (Severe)", "1.00", "1.00", "1.00", "250"],
    ["Weighted average", "0.95", "0.95", "0.95", "1090"],
    ["Overall accuracy", "95.2%", "—", "—", "—"],
]
TABLE_IV = [
    ["Facial region", "Precision", "Recall", "F1-score", "Support"],
    ["Forehead", "0.99", "0.99", "0.99", "93"],
    ["Chin", "0.94", "0.93", "0.93", "329"],
    ["Nose", "0.97", "0.97", "0.97", "103"],
    ["Left cheek", "0.94", "0.93", "0.93", "169"],
    ["Right cheek", "0.96", "0.96", "0.96", "396"],
]

FIGURE_FILES = [
    "fig1_pipeline.jpg",
    "fig2_datasets.jpg",
    "fig3_architecture.jpg",
    "fig4_confusion_en.png",
    "fig5_stability_en.png",
    "fig6_app.jpg",
]

FIGURE_EN_CAPTIONS = [
    "Figure 1 – Overview of the proposed methodology, including image acquisition, preprocessing, Region-Aware ResNet-18 architecture, and Flutter integration.",
    "Figure 2 – Representative Acne04 and Acne Level samples and facial regions used by the proposed model.",
    "Figure 3 – Region-Aware ResNet-18 architecture with learnable region embeddings fused before classification.",
    "Figure 4 – Confusion matrix of the Region-Aware ResNet-18 on the test set.",
    "Figure 5 – Accuracy distribution obtained from 30 independent robustness evaluations.",
    "Figure 6 – Mobile teledermatology prototype developed for automatic facial acne severity assessment.",
]

HEADING_MAP = {
    "Resumo": "Abstract",
    "2. Trabalhos Relacionados": "2. Related Work",
    "2.1 Inteligência Artificial aplicada à Dermatologia": "2.1 Artificial Intelligence Applied to Dermatology",
    "2.2 Classificação Automática da Gravidade da Acne": "2.2 Automatic Acne Severity Classification",
    "2.3 Limitações das Abordagens Atuais": "2.3 Limitations of Current Approaches",
    "3. Materiais e Métodos": "3. Materials and Methods",
    "3.1 Base de Dados e Pré-processamento": "3.1 Dataset and Preprocessing",
    "5. Discussion": "5. Discussion",
    "6. Limitations": "6. Limitations",
    "7. Conclusion": "7. Conclusion",
    "8. Referências": "8. References",
}


def set_paragraph_text(paragraph, text, bold=None, italic=None):
    if not paragraph.runs:
        run = paragraph.add_run(text)
        run.font.name = "Arimo"
        run.font.size = Pt(12)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        return
    paragraph.runs[0].text = text
    if bold is not None:
        paragraph.runs[0].bold = bold
    if italic is not None:
        paragraph.runs[0].italic = italic
    for run in paragraph.runs[1:]:
        run.text = ""


def fill_table(table, rows):
    for r, row in enumerate(rows):
        if r >= len(table.rows):
            break
        for c, val in enumerate(row):
            if c >= len(table.rows[r].cells):
                break
            set_paragraph_text(table.rows[r].cells[c].paragraphs[0], val, bold=(r == 0))


def detect_table_id(rows):
    flat = " ".join(" ".join(r) for r in rows).lower()
    if "característica" in flat or "characteristic" in flat:
        return "I"
    if "parâmetro" in flat or "parameter" in flat or "framework" in flat:
        return "II"
    if "class" in flat and "precision" in flat:
        if "facial region" in flat or "forehead" in flat:
            return "IV"
        return "III"
    return None


def export_pdf(docx_path, pdf_path):
    try:
        import win32com.client

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.ExportAsFixedFormat(
            OutputFileName=os.path.abspath(pdf_path),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            CreateBookmarks=0,
        )
        doc.Close(False)
        word.Quit()
        return os.path.exists(pdf_path)
    except Exception as exc:
        print("win32com PDF failed:", exc)
    try:
        from docx2pdf import convert

        convert(docx_path, pdf_path)
        return os.path.exists(pdf_path)
    except Exception as exc:
        print("docx2pdf failed:", exc)
    return False


def main():
    if not os.path.exists(TEMPLATE):
        print("Template not found:", TEMPLATE)
        sys.exit(1)

    with open(TRANS_JSON, encoding="utf-8") as f:
        translations = json.load(f)

    shutil.copy2(TEMPLATE, OUT_DOCX)
    doc = Document(OUT_DOCX)

    fig_caption_idx = 0
    table_map = {"I": TABLE_I, "II": TABLE_II, "III": TABLE_III, "IV": TABLE_IV}

    for p in doc.paragraphs:
        src = p.text.strip()
        if not src:
            continue

        if src.startswith("Quantidade de Páginas") or src.startswith("Number of Pages"):
            continue

        if src in HEADING_MAP:
            set_paragraph_text(p, HEADING_MAP[src])
            continue

        if src.strip() == "Resumo":
            set_paragraph_text(p, "Abstract")
            continue

        if src.startswith("Palavras-chave:"):
            en = translations.get(src, p.text)
            set_paragraph_text(p, en.replace("Palavras-chave:", "Keywords:"), bold=True)
            continue

        if src.startswith("Figura ") or (src.startswith("Figure ") and "–" in src):
            if fig_caption_idx < len(FIGURE_EN_CAPTIONS):
                set_paragraph_text(p, FIGURE_EN_CAPTIONS[fig_caption_idx], italic=True)
                fig_caption_idx += 1
            continue

        if src in translations:
            set_paragraph_text(p, translations[src])
            continue

        # Partial match for title
        if "Classificação da Gravidade da Acne Facial" in src:
            key = [k for k in translations if "Classificação da Gravidade" in k]
            if key:
                set_paragraph_text(p, translations[key[0]])

    for table in doc.tables:
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        tid = detect_table_id(rows)
        if tid and tid in table_map:
            fill_table(table, table_map[tid])

    # Replace inline images
    fig_paths = [os.path.join(FIGS, f) for f in FIGURE_FILES]
    img_indices = [i for i, p in enumerate(doc.paragraphs) if p._element.xpath(".//a:blip")]
    for idx, fig_path in zip(img_indices[: len(fig_paths)], fig_paths):
        p = doc.paragraphs[idx]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        w = Inches(4.8) if "stability" in fig_path else Inches(5.8)
        run.add_picture(fig_path, width=w)

    # Authors after title
    for i, p in enumerate(doc.paragraphs):
        if "Facial Acne Severity Classification" in p.text or "Region-Aware ResNet-18 with Anatomical" in p.text:
            if i + 1 < len(doc.paragraphs):
                ap = doc.paragraphs[i + 1]
                if len(ap.text.strip()) < 10:
                    author_text = (
                        "Lucas Gonzaga Andrade\nRoger Moura Sarmento\n"
                        "Federal Institute of Education, Science and Technology of Ceará (IFCE) — "
                        "Fortaleza, CE, Brazil"
                    )
                    set_paragraph_text(ap, author_text)
                    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # Insert Table I grid after caption (template has caption only)
    for i, p in enumerate(doc.paragraphs):
        if "Table I –" in p.text or "Tabela I –" in p.text:
            tbl = doc.add_table(rows=len(TABLE_I), cols=len(TABLE_I[0]))
            try:
                tbl.style = "Table Grid"
            except KeyError:
                pass
            fill_table(tbl, TABLE_I)
            p._element.addnext(tbl._tbl)
            body = doc.element.body
            if body[-1] is tbl._tbl:
                body.remove(tbl._tbl)
            break

    # Remove page-count placeholder from template
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("Quantidade de Páginas") or t.startswith("Number of Pages"):
            p._element.getparent().remove(p._element)

    doc.save(OUT_DOCX)
    print("Saved DOCX:", OUT_DOCX)

    if export_pdf(OUT_DOCX, OUT_PDF):
        print("Saved PDF:", OUT_PDF)
    else:
        print("PDF export failed")

    dl = os.path.join(os.path.expanduser("~"), "Downloads")
    shutil.copy2(OUT_DOCX, os.path.join(dl, os.path.basename(OUT_DOCX)))
    if os.path.exists(OUT_PDF):
        shutil.copy2(OUT_PDF, os.path.join(dl, os.path.basename(OUT_PDF)))


if __name__ == "__main__":
    main()
