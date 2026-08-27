# -*- coding: utf-8
"""Condensed English CBEB article (~10 pages) from full translation via paragraph whitelist."""
import json
import os
import re
import shutil

from copy import deepcopy

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_JSON = os.path.join(ROOT, "papers", "cbeb", "source_blocks.json")
TRANS_JSON = os.path.join(ROOT, "papers", "cbeb", "translations_en.json")
FIGS = os.path.join(ROOT, "papers", "sibgrapi", "figs")
OUT_DOCX = os.path.join(ROOT, "papers", "cbeb", "AcneNet_CBEB2026_English_CBEB_Short.docx")
OUT_PDF = os.path.join(ROOT, "papers", "cbeb", "AcneNet_CBEB2026_English_CBEB_Short.pdf")

TEMPLATE_DOCX = r"C:\Users\GonLu\Downloads\[ACNE] - Artigo CBEB 2026 - V-Final.docx"

# Insert numbered equations from the CBEB template after these source paragraph indices.
EQUATIONS_AFTER = {
    56: 1,  # residual block
    62: 3,  # region embedding matrix
    64: 4,  # e_r = E(r)
    67: 5,  # z = Concat(f, e_r)
    72: 6,  # Softmax classifier
}

# Fallback text if template OMML cannot be loaded.
EQUATION_FALLBACK = {
    1: "y = F(x, Wᵢ) + H(x)",
    3: "E ∈ ℝ^{5×64}",
    4: "e_r = E(r)",
    5: "z = Concat(f, e_r) ∈ ℝ^{576}",
    6: "y = Softmax(Wz + b)",
}

# Indices of non-empty paragraphs to KEEP (0-based in filtered paragraph list)
WHITELIST = {
    1, 2, 3, 4, 5, 6, 7, 11, 12, 13,  # title, abstract, intro (skip 8-10, 19)
    20, 21, 22, 25, 26, 28, 30, 34,  # related work (skip 23,24,27,29,31-33)
    35, 36, 39, 40, 41, 42, 43, 46, 47, 48, 50,  # dataset (skip 44 fig ref, 49)
    51, 52, 55, 56, 62, 64, 66, 67, 69, 71, 72,  # arch + training
    74, 75, 76, 77, 78, 79, 80, 81, 84, 85,  # results (skip 82,83,86,87 transitions)
    88, 89,  # robustness (skip 90,112-116 confidence redundancy)
    117, 118, 121, 123,  # discussion (skip 119,120,124)
    125, 126, 129, 130,  # limitations + conclusion (skip 127,128,131-133)
}

TABLE_I = [
    ["Characteristic", "Acne04", "Acne Level"],
    ["Images", "1,406", "1,457"],
    ["Classes", "4 (Hayashi 0–3)", "4 (Hayashi 0–3)"],
    ["Use", "Merged pipeline (70/15/15)", "Merged pipeline (70/15/15)"],
]
TABLE_II = [
    ["Parameter", "Value"],
    ["Backbone", "ResNet-18 (from scratch)"],
    ["Optimizer / loss", "Adam / weighted CE"],
    ["LR / batch / epochs", "1e-4 / 16 / 40"],
    ["Hardware", "RTX 3060 12 GB"],
    ["Weights", "best_robust_model.pth"],
]
TABLE_III = [
    ["Class", "Prec.", "Rec.", "F1", "Sup."],
    ["G0 Clear", "0.96", "0.95", "0.95", "318"],
    ["G1 Mild", "0.94", "0.94", "0.94", "412"],
    ["G2 Moderate", "0.88", "0.92", "0.90", "110"],
    ["G3 Severe", "1.00", "1.00", "1.00", "250"],
    ["W. avg.", "0.95", "0.95", "0.95", "1090"],
]
TABLE_IV = [
    ["Region", "Prec.", "Rec.", "F1", "Sup."],
    ["Forehead", "0.99", "0.99", "0.99", "93"],
    ["Chin", "0.94", "0.93", "0.93", "329"],
    ["Nose", "0.97", "0.97", "0.97", "103"],
    ["L. cheek", "0.94", "0.93", "0.93", "169"],
    ["R. cheek", "0.96", "0.96", "0.96", "396"],
]

FIGURES = [
    ("fig1_pipeline.jpg", "Figure 1 – Proposed methodology from acquisition to mobile reporting.", 4.4),
    ("fig2_datasets.jpg", "Figure 2 – Acne04 / Acne Level samples and extracted facial regions.", 4.4),
    ("fig3_architecture.jpg", "Figure 3 – Region-Aware ResNet-18 with learnable region embeddings.", 4.4),
    ("fig4_confusion_en.png", "Figure 4 – Confusion matrix on the test set.", 4.4),
    ("fig5_stability_en.png", "Figure 5 – Accuracy over 30 robustness runs.", 3.8),
    ("fig6_app.jpg", "Figure 6 – Flutter teledermatology prototype.", 4.2),
]

REFS = [
    "[1] M. Vasam et al., Acne vulgaris review, Biomedicines, 2023.",
    "[2] Y. Li et al., Acne treatment perspectives, Front. Medicine, 2024.",
    "[4] A. H. Heng and F. T. Chew, Acne epidemiology review, Sci. Rep., 2020.",
    "[10] D. Z. Eichenfield et al., Management of acne, JAMA, 2021.",
    "[11] S. Wongvibulsin et al., Dermatology mobile AI apps, JAMA Dermatol., 2024.",
    "[13] S. P. Choy et al., DL for skin disease review, npj Digital Medicine, 2023.",
    "[14] Q. T. Huynh et al., Acne detection and grading, Diagnostics, 2022.",
    "[15] Y. Lin et al., KIEGLFN framework, CMPB, 2022.",
    "[16] S. Liu et al., AcneGrader, Skin Res. Technol., 2022.",
    "[23] K. He et al., ResNet, CVPR, 2016.",
]


def setup_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.54)
        s.right_margin = Cm(2.54)
    st = doc.styles["Normal"]
    st.font.name = "Arimo"
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc


def add_text(doc, text, *, heading=None, bold=False, italic=False, center=False):
    if heading:
        p = doc.add_heading(text, level=heading)
    else:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = "Arimo"
        r.font.size = Pt(12)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Arimo"
    return p


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri].cells[ci].text = val
            for run in t.rows[ri].cells[ci].paragraphs[0].runs:
                run.font.name = "Arimo"
                run.font.size = Pt(10)
                if ri == 0:
                    run.bold = True


def add_figure(doc, fname, caption, width):
    path = os.path.join(FIGS, fname)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
    p = add_text(doc, caption, italic=True, center=True)
    p.runs[0].font.size = Pt(10)


def load_template_equations(template_path):
    """Return {equation_number: paragraph_element} copied from the CBEB template."""
    if not os.path.exists(template_path):
        return {}
    import zipfile
    from xml.etree import ElementTree as ET

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    with zipfile.ZipFile(template_path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    eqs = {}
    for p in root.findall(f".//{{{w_ns}}}p"):
        if p.find(f".//{{{m_ns}}}oMath") is None:
            continue
        label = None
        for t in p.findall(f".//{{{w_ns}}}t"):
            if t.text and re.fullmatch(r"\(\d+\)", t.text.strip()):
                label = int(t.text.strip()[1:-1])
                break
        if label is not None:
            eqs[label] = deepcopy(p)
    return eqs


def add_equation_fallback(doc, eq_text, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(eq_text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(12)
    num = p.add_run(f"\t({number})")
    num.font.name = "Arimo"
    num.font.size = Pt(12)


def insert_body_element(doc, element):
    body = doc.element.body
    if len(body) and body[-1].tag.endswith("sectPr"):
        body.insert(len(body) - 1, element)
    else:
        body.append(element)


def add_equation(doc, number, template_eqs):
    elem = template_eqs.get(number)
    if elem is not None:
        import xml.etree.ElementTree as ET
        from docx.oxml import parse_xml

        insert_body_element(doc, parse_xml(ET.tostring(elem, encoding="unicode")))
    elif number in EQUATION_FALLBACK:
        add_equation_fallback(doc, EQUATION_FALLBACK[number], number)


def export_pdf(docx, pdf):
    try:
        import win32com.client
        w = win32com.client.Dispatch("Word.Application")
        w.Visible = False
        d = w.Documents.Open(os.path.abspath(docx))
        d.ExportAsFixedFormat(os.path.abspath(pdf), 17, OpenAfterExport=False)
        d.Close(False)
        w.Quit()
        return os.path.exists(pdf)
    except Exception as e:
        print("PDF:", e)
        return False


def is_figure_caption(text):
    return text.startswith("Figura ") or (text.startswith("Figure ") and "–" in text)


def is_table_caption(text):
    return text.startswith("Tabela ") or text.startswith("Table ")


def is_equation_label(text):
    return bool(re.match(r"^\s*\(\d+\)\s*$", text.strip()))


def main():
    with open(SRC_JSON, encoding="utf-8") as f:
        blocks = json.load(f)
    with open(TRANS_JSON, encoding="utf-8") as f:
        trans = json.load(f)

    paras = [b for b in blocks if b["type"] == "p" and b["text"].strip()]
    doc = setup_doc()
    template_eqs = load_template_equations(TEMPLATE_DOCX)
    fig_i = 0
    pending_figure = None
    paragraphs_since_table = 99

    def note_body_paragraph():
        nonlocal paragraphs_since_table, pending_figure
        paragraphs_since_table += 1
        if pending_figure and paragraphs_since_table >= 1:
            fname, cap, w = pending_figure
            add_figure(doc, fname, cap, w)
            pending_figure = None
            paragraphs_since_table = 0

    def queue_or_add_figure():
        nonlocal fig_i, pending_figure, paragraphs_since_table
        if fig_i >= len(FIGURES):
            return
        item = FIGURES[fig_i]
        fig_i += 1
        if paragraphs_since_table < 1:
            pending_figure = item
        else:
            fname, cap, w = item
            add_figure(doc, fname, cap, w)
            paragraphs_since_table = 0

    def maybe_insert_equation(idx):
        eq_num = EQUATIONS_AFTER.get(idx)
        if eq_num is not None:
            add_equation(doc, eq_num, template_eqs)

    for idx, block in enumerate(paras):
        src = block["text"]
        if idx not in WHITELIST and not is_figure_caption(src):
            continue
        if src.startswith("Quantidade de Páginas"):
            continue
        if is_equation_label(src):
            continue

        en = trans.get(src, src)

        if "Classificação da Gravidade" in src:
            add_text(doc, en, center=True, bold=True)
            add_text(
                doc,
                "Lucas Gonzaga Andrade\nRoger Moura Sarmento\n"
                "IFCE — Fortaleza, CE, Brazil",
                center=True,
            )
            continue

        if src.strip() == "Resumo":
            add_text(doc, "Abstract", heading=1)
            continue

        if src.startswith("Palavras-chave"):
            add_text(doc, en.replace("Palavras-chave:", "Keywords:"), bold=True)
            continue

        if src == "As principais contribuições deste trabalho são:":
            add_text(
                doc,
                "Contributions: Region-Aware ResNet-18 with 64-D region embeddings; "
                "YOLOv8-based five-zone preprocessing; extended evaluation (global, regional, "
                "stability, confidence); Flutter teledermatology prototype.",
            )
            continue

        if is_table_caption(src):
            add_text(doc, en, italic=True, center=True)
            if "Table I" in en or "Tabela I" in src:
                add_table(doc, TABLE_I)
            elif "Table II" in en or "Tabela II" in src:
                add_table(doc, TABLE_II)
            elif "Table III" in en or "Tabela III" in src:
                add_table(doc, TABLE_III)
            elif "Table IV" in en or "Tabela IV" in src:
                add_table(doc, TABLE_IV)
            paragraphs_since_table = 0
            continue

        if is_figure_caption(src):
            queue_or_add_figure()
            continue

        if block.get("style", "").startswith("Heading") or re.match(r"^\d+\.", src.strip()):
            lvl = 2 if re.match(r"^\d+\.\d+", src.strip()) else 1
            add_text(doc, en, heading=lvl)
            note_body_paragraph()
            maybe_insert_equation(idx)
            continue

        if src.strip() == "8. Referências" or src.strip().startswith("[1]"):
            continue

        add_text(doc, en)
        note_body_paragraph()
        maybe_insert_equation(idx)

    if pending_figure:
        fname, cap, w = pending_figure
        add_figure(doc, fname, cap, w)

    add_text(doc, "References", heading=1)
    for r in REFS:
        add_text(doc, r)

    doc.save(OUT_DOCX)

    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(OUT_DOCX) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    words = len(re.findall(r"\S+", " ".join(t.text for t in root.findall(".//w:t", ns) if t.text)))
    print(f"Saved DOCX: {OUT_DOCX} (~{words} words)")

    if export_pdf(OUT_DOCX, OUT_PDF):
        print(f"Saved PDF: {OUT_PDF}")

    dl = os.path.join(os.path.expanduser("~"), "Downloads")
    shutil.copy2(OUT_DOCX, os.path.join(dl, os.path.basename(OUT_DOCX)))
    if os.path.exists(OUT_PDF):
        shutil.copy2(OUT_PDF, os.path.join(dl, os.path.basename(OUT_PDF)))


if __name__ == "__main__":
    main()
