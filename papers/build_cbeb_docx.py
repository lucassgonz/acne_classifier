# -*- coding: utf-8 -*-
"""Generate condensed CBEB 2026 article (4-8 pages) as DOCX."""
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

ROOT = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(ROOT, "figs")
OUT = os.path.join(ROOT, "cbeb", "AcneNet_CBEB2026_Curto.docx")
os.makedirs(os.path.join(ROOT, "cbeb"), exist_ok=True)

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def add_authors(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)


def add_heading(text, level=1):
    doc.add_heading(text, level=level)


def add_para(text):
    doc.add_paragraph(text)


def add_figure(filename, caption, width=Inches(6.2)):
    path = os.path.join(FIGS, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=width)
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(10)


add_title(
    "Classificacao da Gravidade da Acne Facial Utilizando "
    "ResNet-18 Sensivel a Regiao com Embeddings Anatomicos"
)
add_authors(
    "Lucas Gonzaga Andrade\nRoger Moura Sarmento\n"
    "Instituto Federal de Educacao, Ciencia e Tecnologia do Ceara (IFCE) — Fortaleza, CE"
)

add_heading("Resumo", level=1)
add_para(
    "A acne vulgar e uma das dermatoses mais prevalentes, e a avaliacao de sua gravidade "
    "continua dependente de inspecao clinica subjetiva. Este trabalho propoe uma arquitetura "
    "Region-Aware ResNet-18 para classificacao automatica em quatro niveis de gravidade a "
    "partir de recortes de regioes faciais. Regioes sao extraidas com YOLOv8 e classificadas "
    "por um unico backbone ResNet-18 enriquecido com embeddings aprendiveis (testa, queixo, "
    "nariz e bochechas). O modelo foi treinado com dados combinados dos corpora Acne04 e "
    "Acne Level, utilizando aumento de dados e cross-entropy ponderada. Em 1.090 recortes de "
    "teste, obteve-se 95,2% de acuracia e F1-score ponderado de 0,95, com desempenho "
    "consistente entre regioes (F1 >= 0,93). Um prototipo Flutter ilustra integracao em "
    "teledermatologia, embora a inferencia em dispositivo ainda nao esteja implementada."
)
add_para("Palavras-chave: acne vulgar; aprendizado profundo; ResNet-18; teledermatologia; visao computacional.")

add_heading("1. Introducao", level=1)
add_para(
    "A avaliacao da gravidade da acne exige experiencia clinica e apresenta variabilidade "
    "inter-observador. Solucoes baseadas em multiplos estagios — deteccao de lesoes seguida "
    "de classificacao — aumentam custo computacional e propagam erros. Propomos um classificador "
    "unificado que recebe recortes de regiao facial e um identificador anatomico, preservando "
    "contexto regional com baixa complexidade."
)

add_heading("2. Materiais e Metodos", level=1)
add_heading("2.1 Dados e pre-processamento", level=2)
add_para(
    "Utilizamos dois corpora publicos: Acne04 (acne_1024, 1.406 imagens) e Acne Level "
    "(Dataset/, 1.457 imagens), com quatro niveis de gravidade (0–3). As imagens foram "
    "consolidadas, segmentadas com YOLOv8 em cinco regioes faciais e divididas em "
    "treino/validacao/teste (70/15/15). Cada regiao detectada foi redimensionada para 224x224."
)
add_figure("from_docx/image2.png", "Figura 1 – Visao geral do pipeline: aquisicao, segmentacao YOLOv8, classificacao e prototipo movel.")

add_heading("2.2 Arquitetura Region-Aware ResNet-18", level=2)
add_para(
    "A ResNet-18 extrai features visuais (512 dimensoes), concatenadas a um embedding "
    "aprendivel de 64 dimensoes da regiao facial. Um classificador fully connected com "
    "BatchNorm e Dropout produz quatro logits de gravidade. O backbone foi treinado do "
    "zero (sem pesos ImageNet)."
)
add_figure("from_docx/image3.png", "Figura 2 – Arquitetura Region-Aware ResNet-18.", width=Inches(5.8))

add_heading("2.3 Treinamento", level=2)
add_para(
    "Framework PyTorch; otimizador Adam (lr = 0,0001); cross-entropy ponderada; batch 16; "
    "40 epocas maximas; ReduceLROnPlateau; early stopping. Augmentations: flip horizontal, "
    "rotacao (±25°), color jitter, blur gaussiano e random erasing."
)

add_heading("3. Resultados", level=1)
add_para(
    "Tabela 1 resume as metricas globais no conjunto de teste (1.090 recortes). A acuracia "
    "foi de 95,2% e o F1-score ponderado de 0,95. A maioria dos erros ocorreu entre classes "
    "adjacentes (1 e 2). Por regiao, os F1-scores ponderados foram: testa 0,99; queixo 0,93; "
    "nariz 0,97; bochecha esquerda 0,93; bochecha direita 0,96."
)

table = doc.add_table(rows=6, cols=5)
table.style = "Table Grid"
headers = ["Classe", "Precision", "Recall", "F1", "Support"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
rows = [
    ["0", "0,96", "0,95", "0,95", "318"],
    ["1", "0,94", "0,94", "0,94", "412"],
    ["2", "0,88", "0,92", "0,90", "110"],
    ["3", "1,00", "1,00", "1,00", "250"],
    ["Media pond.", "0,95", "0,95", "0,95", "1090"],
]
for r, row in enumerate(rows, start=1):
    for c, val in enumerate(row):
        table.rows[r].cells[c].text = val
doc.add_paragraph("Tabela 1 – Desempenho global no conjunto de teste.").alignment = WD_ALIGN_PARAGRAPH.CENTER

add_figure(
    "figure4_confusion_matrix_en.png",
    "Figura 3 – Matriz de confusao: contagens absolutas (esquerda) e percentuais normalizados (direita). Labels em ingles conforme solicitacao.",
)

add_para(
    "Analise de estabilidade com 30 execucoes e perturbacoes leves reportou acuracia media "
    "de 94,41% (desvio-padrao 0,46%), indicando robustez operacional."
)
add_figure("from_docx/image5.png", "Figura 4 – Distribuicao da acuracia em 30 avaliacoes de robustez.")

add_heading("4. Prototipo movel", level=1)
add_para(
    "Desenvolvemos um aplicativo Flutter para captura de imagens e fluxo de teledermatologia. "
    "A inferencia automatica ainda nao esta integrada; o prototipo demonstra viabilidade de "
    "interface, nao validacao clinica."
)
add_figure("from_docx/image6.png", "Figura 5 – Prototipo movel de teledermatologia.", width=Inches(5.5))

add_heading("5. Conclusao", level=1)
add_para(
    "A arquitetura Region-Aware ResNet-18 classificou a gravidade da acne com 95,2% de "
    "acuracia em recortes regionais, mantendo desempenho estavel entre regioes faciais. "
    "Trabalhos futuros incluem calibracao de confianca, validacao externa, inferencia em "
    "dispositivo movel e maior diversidade de fototipos."
)

add_heading("Referencias", level=1)
refs = [
    "[1] S. P. Choy et al., Systematic review of deep learning image analyses for skin disease, npj Digital Medicine, 2023.",
    "[2] Q. T. Huynh et al., Automatic acne detection and severity grading, Diagnostics, 2022.",
    "[3] Y. Lin et al., KIEGLFN: unified acne grading framework, CMPB, 2022.",
    "[4] K. He et al., Deep residual learning for image recognition, CVPR, 2016.",
    "[5] N. Hayashi et al., Acne severity grading system, Journal of Dermatology, 2008.",
]
for r in refs:
    add_para(r)

doc.save(OUT)
print("Saved", OUT)
